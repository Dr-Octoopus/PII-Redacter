"""
Parser for .xlsx (Excel) files.

Design choice: spreadsheets don't have "paragraphs" -- they have
cells. So instead of paragraph offsets, we track (sheet_name, cell_
coordinate) for every non-empty cell. This lets the redactor go
back and overwrite the exact cell where PII was found, leaving
every other cell, formula, and formatting untouched.
"""

from openpyxl import load_workbook


def parse_xlsx(filepath: str) -> tuple[str, list[dict]]:
    """
    Extract text from all cells across all sheets in an .xlsx file.

    Args:
        filepath: Path to the .xlsx file on disk.

    Returns:
        A tuple of:
          - full_text: every non-empty cell's value joined with
            newlines (used for running detection).
          - cell_map: a list of dicts, one per non-empty cell,
            containing sheet name, cell coordinate (e.g. "B7"), the
            cell's text, and its (start, end) offsets within
            full_text -- used to trace a detected entity back to
            its exact cell.
    """
    workbook = load_workbook(filepath)

    full_text_parts = []
    cell_map = []
    running_offset = 0

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue

                text = str(cell.value)
                start = running_offset
                end = start + len(text)

                cell_map.append({
                    "sheet_name": sheet.title,
                    "coordinate": cell.coordinate,
                    "text": text,
                    "start": start,
                    "end": end,
                })

                full_text_parts.append(text)
                running_offset = end + 1

    full_text = "\n".join(full_text_parts)
    return full_text, cell_map
